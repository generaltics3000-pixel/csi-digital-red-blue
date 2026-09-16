import pptx

s3_img_path = "/home/dario/.gemini/antigravity-ide/brain/b9af9946-92a1-4c0c-9a97-0b6a6e1b42f5/scratch/slide3_rounded.png"
s7_img_path = "/home/dario/.gemini/antigravity-ide/brain/b9af9946-92a1-4c0c-9a97-0b6a6e1b42f5/scratch/slide7_rounded.png"

with open(s3_img_path, "rb") as f:
    s3_blob = f.read()

with open(s7_img_path, "rb") as f:
    s7_blob = f.read()

files_to_update = [
    "Ciudadanía Digital Segura.pptx",
    "Ciudadanía Digital Segura - Taller Individual WhatsApp.pptx"
]

for filename in files_to_update:
    prs = pptx.Presentation(filename)
    
    # 1. Slide 3 (index 2) - Phishing image
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE and shape.width > 5000000 and shape.left > 1000000:
            blip = shape._element.xpath('.//a:blip')
            if blip:
                rId = blip[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                part = shape.part.related_part(rId)
                # We update the blob of rId5 (which is 8294 bytes placeholder)
                if len(part.blob) < 20000:
                    print(f"Updating Slide 3 image in {filename} (shape {shape.name}, rId {rId}, old size {len(part.blob)})")
                    part._blob = s3_blob
                    
    # 2. Slide 7 (index 6) - VPN image
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE and shape.width > 5000000 and shape.left > 1000000:
            blip = shape._element.xpath('.//a:blip')
            if blip:
                rId = blip[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                part = shape.part.related_part(rId)
                if len(part.blob) < 20000:
                    print(f"Updating Slide 7 image in {filename} (shape {shape.name}, rId {rId}, old size {len(part.blob)})")
                    part._blob = s7_blob

    prs.save(filename)
    print(f"✓ Saved updated {filename}")

