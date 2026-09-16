import os
import io
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_presentation():
    src_path = "Ciudadanía Digital Segura_backup.pptx"
    prs = pptx.Presentation(src_path)

    bg_blob = prs.slides[0].shapes[0].image.blob
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 10: Anatomía del Engaño en WhatsApp
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    slide1.shapes.add_picture(io.BytesIO(bg_blob), 0, 0, Inches(13.333), Inches(7.50))

    tb_title1 = slide1.shapes.add_textbox(Inches(0.625), Inches(0.48), Inches(12.08), Inches(0.55))
    tf1 = tb_title1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0
    p1 = tf1.paragraphs[0]
    p1.text = "Anatomía del Engaño en WhatsApp"
    p1.font.name = "Montserrat"
    p1.font.size = Pt(27)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(248, 250, 252)

    tb_sub1 = slide1.shapes.add_textbox(Inches(0.625), Inches(1.12), Inches(12.08), Inches(0.45))
    tfs1 = tb_sub1.text_frame
    tfs1.word_wrap = True
    tfs1.margin_left = tfs1.margin_top = tfs1.margin_right = tfs1.margin_bottom = 0
    ps1 = tfs1.paragraphs[0]
    
    r1 = ps1.add_run()
    r1.text = "Desglosamos la estructura de un mensaje fraudulento (Smishing): "
    r1.font.name = "Inter"
    r1.font.size = Pt(13)
    r1.font.color.rgb = RGBColor(148, 163, 184)
    
    r2 = ps1.add_run()
    r2.text = "cómo operan los ciberatacantes paso a paso"
    r2.font.name = "Inter"
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(56, 189, 248)

    r3 = ps1.add_run()
    r3.text = "."
    r3.font.name = "Inter"
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGBColor(148, 163, 184)

    # 4 Cards Layout
    c4_w = Inches(2.80)
    c4_h = Inches(5.30)
    c4_y = Inches(1.70)
    c4_gap = Inches(0.29)

    cards_data_10 = [
        {
            "border": RGBColor(14, 165, 233),
            "badge_bg": RGBColor(8, 47, 73),
            "badge_txt": "1. SUPLANTACIÓN",
            "badge_col": RGBColor(186, 230, 253),
            "title": "Falso Remitente",
            "title_col": RGBColor(56, 189, 248),
            "subtitle": "Generar confianza inmediata.",
            "items": [
                ("Logo Robado: ", "Uso de marcas conocidas (WhatsApp, MercadoPago, Bancos)."),
                ("Nombre Formal: ", "'Soporte Oficial' o 'Beca Conectividad 2026'."),
                ("Número Raro: ", "Prefijos internacionales o números no agendados.")
            ],
            "callout_bg": RGBColor(12, 26, 46),
            "callout_border": RGBColor(3, 105, 161),
            "callout_txt": "🎯 Falsa autoridad visual"
        },
        {
            "border": RGBColor(245, 158, 11),
            "badge_bg": RGBColor(120, 53, 15),
            "badge_txt": "2. EL GANCHO",
            "badge_col": RGBColor(253, 230, 138),
            "title": "Disparador Emocional",
            "title_col": RGBColor(251, 191, 36),
            "subtitle": "Bloquear el pensamiento crítico.",
            "items": [
                ("Urgencia Extrema: ", "'Tu cuenta se bloqueará en 30 minutos'."),
                ("Premio / Subsidio: ", "Falsas promesas de dinero fácil o regalos."),
                ("Curiosidad / Miedo: ", "Alerta sobre supuestas compras no autorizadas.")
            ],
            "callout_bg": RGBColor(32, 24, 12),
            "callout_border": RGBColor(180, 83, 9),
            "callout_txt": "⚡ Forzar acción impulsiva"
        },
        {
            "border": RGBColor(239, 68, 68),
            "badge_bg": RGBColor(69, 10, 10),
            "badge_txt": "3. LA TRAMPA",
            "badge_col": RGBColor(254, 202, 202),
            "title": "Trampa Técnica",
            "title_col": RGBColor(248, 113, 113),
            "subtitle": "El vector del ciberataque.",
            "items": [
                ("Enlace Falso: ", "URLs engañosas (.free.nf, banconac1on.xyz)."),
                ("Robo de Código SMS: ", "Piden el código de 6 dígitos para robar WhatsApp."),
                ("Adjuntos Trampa: ", "Archivos ejecutables (.apk/.exe) disfrazados.")
            ],
            "callout_bg": RGBColor(24, 18, 30),
            "callout_border": RGBColor(153, 27, 27),
            "callout_txt": "🪤 Robo de credenciales"
        },
        {
            "border": RGBColor(139, 92, 246),
            "badge_bg": RGBColor(76, 29, 149),
            "badge_txt": "4. AISLAMIENTO",
            "badge_col": RGBColor(221, 214, 254),
            "title": "Presión & Silencio",
            "title_col": RGBColor(196, 181, 253),
            "subtitle": "Evitar que consultes con otros.",
            "items": [
                ("Confidencialidad: ", "'No comentes con nadie por seguridad bancaria'."),
                ("Sin Canales Reales: ", "Exigen no llamar al 0800 ni ir a la sucursal."),
                ("Acoso Digital: ", "Insistencia para no dar tiempo a reflexionar.")
            ],
            "callout_bg": RGBColor(24, 16, 42),
            "callout_border": RGBColor(109, 40, 217),
            "callout_txt": "🛑 Bloquear la verificación"
        }
    ]

    for i, c in enumerate(cards_data_10):
        left_pos = Inches(0.625) + i * (c4_w + c4_gap)

        # Card shape
        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, c4_y, c4_w, c4_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 23, 42)
        card.line.color.rgb = c["border"]
        card.line.width = Pt(1.5)

        # Badge
        badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.18), Inches(1.98), c4_w - Inches(0.36), Inches(0.34))
        badge.fill.solid()
        badge.fill.fore_color.rgb = c["badge_bg"]
        badge.line.color.rgb = c["border"]
        badge.line.width = Pt(1)
        btf = badge.text_frame
        btf.word_wrap = False
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run()
        br.text = c["badge_txt"]
        br.font.name = "Montserrat"
        br.font.size = Pt(9.5)
        br.font.bold = True
        br.font.color.rgb = c["badge_col"]

        # Content
        tb = slide1.shapes.add_textbox(left_pos + Inches(0.18), Inches(2.46), c4_w - Inches(0.36), Inches(3.60))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        ph = tf.paragraphs[0]
        ph.text = c["title"]
        ph.font.name = "Montserrat"
        ph.font.size = Pt(15)
        ph.font.bold = True
        ph.font.color.rgb = c["title_col"]
        ph.space_after = Pt(2)

        psub = tf.add_paragraph()
        psub.text = c["subtitle"]
        psub.font.name = "Inter"
        psub.font.size = Pt(10)
        psub.font.color.rgb = RGBColor(148, 163, 184)
        psub.space_after = Pt(10)

        for ititle, ibody in c["items"]:
            pi = tf.add_paragraph()
            pi.space_after = Pt(7)
            
            rt = pi.add_run()
            rt.text = "• " + ititle
            rt.font.name = "Inter"
            rt.font.size = Pt(10.5)
            rt.font.bold = True
            rt.font.color.rgb = RGBColor(241, 245, 249)

            rb = pi.add_run()
            rb.text = ibody
            rb.font.name = "Inter"
            rb.font.size = Pt(10)
            rb.font.color.rgb = RGBColor(148, 163, 184)

        # Callout Pill
        callout = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos + Inches(0.18), Inches(6.20), c4_w - Inches(0.36), Inches(0.56))
        callout.fill.solid()
        callout.fill.fore_color.rgb = c["callout_bg"]
        callout.line.color.rgb = c["callout_border"]
        callout.line.width = Pt(1)
        ctf = callout.text_frame
        ctf.word_wrap = True
        ctf.margin_top = Inches(0.08)
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = c["callout_txt"]
        cr.font.name = "Inter"
        cr.font.size = Pt(9.5)
        cr.font.bold = True
        cr.font.color.rgb = c["badge_col"]


    # =========================================================================
    # SLIDE 11: Señales de Alerta: Análisis Forense de un Chat
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.shapes.add_picture(io.BytesIO(bg_blob), 0, 0, Inches(13.333), Inches(7.50))

    tb_title2 = slide2.shapes.add_textbox(Inches(0.625), Inches(0.48), Inches(12.08), Inches(0.55))
    tf2 = tb_title2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = "Señales de Alerta: Auditoría de un Chat"
    p2.font.name = "Montserrat"
    p2.font.size = Pt(27)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(248, 250, 252)

    tb_sub2 = slide2.shapes.add_textbox(Inches(0.625), Inches(1.12), Inches(12.08), Inches(0.45))
    tfs2 = tb_sub2.text_frame
    tfs2.word_wrap = True
    tfs2.margin_left = tfs2.margin_top = tfs2.margin_right = tfs2.margin_bottom = 0
    ps2 = tfs2.paragraphs[0]
    
    r2_1 = ps2.add_run()
    r2_1.text = "Tres filtros forenses para "
    r2_1.font.name = "Inter"
    r2_1.font.size = Pt(13)
    r2_1.font.color.rgb = RGBColor(148, 163, 184)
    
    r2_2 = ps2.add_run()
    r2_2.text = "desarmar cualquier engaño"
    r2_2.font.name = "Inter"
    r2_2.font.size = Pt(13)
    r2_2.font.bold = True
    r2_2.font.color.rgb = RGBColor(56, 189, 248)

    r2_3 = ps2.add_run()
    r2_3.text = " antes de hacer clic o responder."
    r2_3.font.name = "Inter"
    r2_3.font.size = Pt(13)
    r2_3.font.color.rgb = RGBColor(148, 163, 184)

    # 3 Cards Layout
    c3_w = Inches(3.82)
    c3_h = Inches(5.30)
    c3_y = Inches(1.70)

    cards_data_11 = [
        {
            "left": Inches(0.625),
            "border": RGBColor(14, 165, 233),
            "badge_bg": RGBColor(8, 47, 73),
            "badge_txt": "FILTRO 1",
            "badge_col": RGBColor(186, 230, 253),
            "title": "1. Auditar el Remitente",
            "title_col": RGBColor(56, 189, 248),
            "subtitle": "Inspeccionar el origen del contacto.",
            "items": [
                ("🚩 Número Desconocido: ", "Si no está agendado o usa un prefijo internacional (+62, +234, +1), sospecha."),
                ("🚩 Falta de Tilde Verde: ", "Las empresas oficiales poseen verificación de WhatsApp certificada."),
                ("🚩 Foto de Perfil: ", "Logos pixelados o descargados de internet son indicio claro de fraude.")
            ],
            "callout_bg": RGBColor(12, 26, 46),
            "callout_border": RGBColor(3, 105, 161),
            "callout_txt": "🔍 Verificar en canales oficiales"
        },
        {
            "left": Inches(4.756),
            "border": RGBColor(245, 158, 11),
            "badge_bg": RGBColor(120, 53, 15),
            "badge_txt": "FILTRO 2",
            "badge_col": RGBColor(253, 230, 138),
            "title": "2. Inspeccionar el Enlace",
            "title_col": RGBColor(251, 191, 36),
            "subtitle": "Examinar minuciosamente la URL.",
            "items": [
                ("🚩 Dominios Falsos: ", "Revisa la dirección exacta (.xyz, .free.nf, .top en vez de .com o .gob.ar)."),
                ("🚩 Typosquatting: ", "Letras cambiadas que imitan sitios reales (ej. mercadopag0.com)."),
                ("🚩 Acortadores Ocultos: ", "Enlaces tipo bit.ly o tinyurl que esconden la web de destino.")
            ],
            "callout_bg": RGBColor(32, 24, 12),
            "callout_border": RGBColor(180, 83, 9),
            "callout_txt": "🚫 No ingresar contraseñas en links"
        },
        {
            "left": Inches(8.888),
            "border": RGBColor(16, 185, 129),
            "badge_bg": RGBColor(6, 78, 59),
            "badge_txt": "FILTRO 3",
            "badge_col": RGBColor(167, 243, 208),
            "title": "3. Regla del Código SMS",
            "title_col": RGBColor(52, 211, 153),
            "subtitle": "La última barrera de defensa.",
            "items": [
                ("🚩 Código de 6 Dígitos: ", "Si te piden el código SMS que te acaba de llegar, quieren robar tu WhatsApp."),
                ("🚩 El Banco Nunca Pide Claves: ", "Ninguna entidad solicita contraseñas o Tokens por mensaje."),
                ("🚩 Regla de los 3 Segundos: ", "Detente, piensa con calma y no actúes impulsivamente bajo presión.")
            ],
            "callout_bg": RGBColor(10, 32, 24),
            "callout_border": RGBColor(4, 120, 87),
            "callout_txt": "🛡️ Tu código SMS es confidencial"
        }
    ]

    for c in cards_data_11:
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c["left"], c3_y, c3_w, c3_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 23, 42)
        card.line.color.rgb = c["border"]
        card.line.width = Pt(1.5)

        badge = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c["left"] + Inches(0.25), Inches(1.98), Inches(1.3), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = c["badge_bg"]
        badge.line.color.rgb = c["border"]
        badge.line.width = Pt(1)
        btf = badge.text_frame
        btf.word_wrap = False
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        br = bp.add_run()
        br.text = c["badge_txt"]
        br.font.name = "Montserrat"
        br.font.size = Pt(10)
        br.font.bold = True
        br.font.color.rgb = c["badge_col"]

        tb = slide2.shapes.add_textbox(c["left"] + Inches(0.25), Inches(2.46), c3_w - Inches(0.5), Inches(3.60))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        ph = tf.paragraphs[0]
        ph.text = c["title"]
        ph.font.name = "Montserrat"
        ph.font.size = Pt(16)
        ph.font.bold = True
        ph.font.color.rgb = c["title_col"]
        ph.space_after = Pt(3)

        psub = tf.add_paragraph()
        psub.text = c["subtitle"]
        psub.font.name = "Inter"
        psub.font.size = Pt(10.5)
        psub.font.color.rgb = RGBColor(148, 163, 184)
        psub.space_after = Pt(12)

        for ititle, ibody in c["items"]:
            pi = tf.add_paragraph()
            pi.space_after = Pt(8)
            
            rt = pi.add_run()
            rt.text = ititle
            rt.font.name = "Inter"
            rt.font.size = Pt(11)
            rt.font.bold = True
            rt.font.color.rgb = RGBColor(241, 245, 249)

            rb = pi.add_run()
            rb.text = ibody
            rb.font.name = "Inter"
            rb.font.size = Pt(10.5)
            rb.font.color.rgb = RGBColor(148, 163, 184)

        callout = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c["left"] + Inches(0.25), Inches(6.20), c3_w - Inches(0.5), Inches(0.56))
        callout.fill.solid()
        callout.fill.fore_color.rgb = c["callout_bg"]
        callout.line.color.rgb = c["callout_border"]
        callout.line.width = Pt(1)
        ctf = callout.text_frame
        ctf.word_wrap = True
        ctf.margin_top = Inches(0.08)
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = c["callout_txt"]
        cr.font.name = "Inter"
        cr.font.size = Pt(10)
        cr.font.bold = True
        cr.font.color.rgb = c["badge_col"]


    # =========================================================================
    # SLIDE 12: Actividad Práctica: Laboratorio de Phishing (Modo Armador)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    slide3.shapes.add_picture(io.BytesIO(bg_blob), 0, 0, Inches(13.333), Inches(7.50))

    tb_title3 = slide3.shapes.add_textbox(Inches(0.625), Inches(0.48), Inches(12.08), Inches(0.55))
    tf3 = tb_title3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0
    p3 = tf3.paragraphs[0]
    p3.text = "Actividad Práctica: Modo Armador"
    p3.font.name = "Montserrat"
    p3.font.size = Pt(27)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(248, 250, 252)

    tb_sub3 = slide3.shapes.add_textbox(Inches(0.625), Inches(1.12), Inches(12.08), Inches(0.45))
    tfs3 = tb_sub3.text_frame
    tfs3.word_wrap = True
    tfs3.margin_left = tfs3.margin_top = tfs3.margin_right = tfs3.margin_bottom = 0
    ps3 = tfs3.paragraphs[0]
    
    r3_1 = ps3.add_run()
    r3_1.text = "Misión individual: ponte en el rol de creador para "
    r3_1.font.name = "Inter"
    r3_1.font.size = Pt(13)
    r3_1.font.color.rgb = RGBColor(148, 163, 184)
    
    r3_2 = ps3.add_run()
    r3_2.text = "comprender la ingeniería social desde adentro"
    r3_2.font.name = "Inter"
    r3_2.font.size = Pt(13)
    r3_2.font.bold = True
    r3_2.font.color.rgb = RGBColor(56, 189, 248)

    r3_3 = ps3.add_run()
    r3_3.text = "."
    r3_3.font.name = "Inter"
    r3_3.font.size = Pt(13)
    r3_3.font.color.rgb = RGBColor(148, 163, 184)

    card_w2 = Inches(5.88)
    card_h2 = Inches(5.30)
    card_y2 = Inches(1.70)

    # Left Card: Consigna Individual
    c12_left = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.625), card_y2, card_w2, card_h2)
    c12_left.fill.solid()
    c12_left.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c12_left.line.color.rgb = RGBColor(236, 72, 153)
    c12_left.line.width = Pt(1.5)

    b12_left = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.98), Inches(3.2), Inches(0.36))
    b12_left.fill.solid()
    b12_left.fill.fore_color.rgb = RGBColor(80, 7, 36)
    b12_left.line.color.rgb = RGBColor(236, 72, 153)
    b12_left.line.width = Pt(1)
    b12l_tf = b12_left.text_frame
    b12l_tf.word_wrap = False
    b12l_p = b12l_tf.paragraphs[0]
    b12l_p.alignment = PP_ALIGN.CENTER
    b12l_r = b12l_p.add_run()
    b12l_r.text = "🎯 MISIÓN INDIVIDUAL"
    b12l_r.font.name = "Montserrat"
    b12l_r.font.size = Pt(11)
    b12l_r.font.bold = True
    b12l_r.font.color.rgb = RGBColor(251, 207, 232)

    tb12_l = slide3.shapes.add_textbox(Inches(0.95), Inches(2.46), Inches(5.23), Inches(3.60))
    tf12_l = tb12_l.text_frame
    tf12_l.word_wrap = True
    tf12_l.margin_left = tf12_l.margin_top = tf12_l.margin_right = tf12_l.margin_bottom = 0

    p12l_h = tf12_l.paragraphs[0]
    p12l_h.text = "Armá tu Mensaje de Engaño"
    p12l_h.font.name = "Montserrat"
    p12l_h.font.size = Pt(17)
    p12l_h.font.bold = True
    p12l_h.font.color.rgb = RGBColor(244, 114, 182)
    p12l_h.space_after = Pt(3)

    p12l_sub = tf12_l.add_paragraph()
    p12l_sub.text = "Elegirás los 4 componentes según el objetivo secreto asignado por el facilitador."
    p12l_sub.font.name = "Inter"
    p12l_sub.font.size = Pt(11)
    p12l_sub.font.color.rgb = RGBColor(148, 163, 184)
    p12l_sub.space_after = Pt(12)

    items_12l = [
        ("🎯 1. Objetivo Asignado: ", "Robar un código de verificación, clonar credenciales o simular una beca."),
        ("📱 2. Remitente: ", "Elegir entre amigo guardado, soporte oficial de marca o número desconocido."),
        ("⚡ 3. Gancho Psicológico: ", "Seleccionar el motivo: urgencia extrema, premio en efectivo o bloqueo."),
        ("🪤 4. Trampa y Cierre: ", "Añadir enlace fraudulento, pedido de código o advertencia de confidencialidad.")
    ]

    for title_txt, body_txt in items_12l:
        pi = tf12_l.add_paragraph()
        pi.space_after = Pt(8)
        
        rt = pi.add_run()
        rt.text = title_txt
        rt.font.name = "Inter"
        rt.font.size = Pt(11.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(253, 231, 243)

        rb = pi.add_run()
        rb.text = body_txt
        rb.font.name = "Inter"
        rb.font.size = Pt(11)
        rb.font.color.rgb = RGBColor(148, 163, 184)

    callout12_l = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(6.20), Inches(5.23), Inches(0.56))
    callout12_l.fill.solid()
    callout12_l.fill.fore_color.rgb = RGBColor(30, 14, 24)
    callout12_l.line.color.rgb = RGBColor(190, 24, 93)
    callout12_l.line.width = Pt(1)
    c12l_tf = callout12_l.text_frame
    c12l_tf.word_wrap = True
    c12l_tf.margin_top = Inches(0.08)
    c12l_p = c12l_tf.paragraphs[0]
    c12l_p.alignment = PP_ALIGN.CENTER
    c12l_r1 = c12l_p.add_run()
    c12l_r1.text = "📱 En tu Celular: "
    c12l_r1.font.name = "Inter"
    c12l_r1.font.size = Pt(10.5)
    c12l_r1.font.bold = True
    c12l_r1.font.color.rgb = RGBColor(244, 114, 182)
    c12l_r2 = c12l_p.add_run()
    c12l_r2.text = "Seleccionar las 4 opciones y enviar tu mensaje."
    c12l_r2.font.name = "Inter"
    c12l_r2.font.size = Pt(10)
    c12l_r2.font.color.rgb = RGBColor(226, 232, 240)


    # Right Card: Simulador en Vivo
    c12_right = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.828), card_y2, card_w2, card_h2)
    c12_right.fill.solid()
    c12_right.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c12_right.line.color.rgb = RGBColor(14, 165, 233)
    c12_right.line.width = Pt(1.5)

    b12_right = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.153), Inches(1.98), Inches(3.2), Inches(0.36))
    b12_right.fill.solid()
    b12_right.fill.fore_color.rgb = RGBColor(8, 47, 73)
    b12_right.line.color.rgb = RGBColor(14, 165, 233)
    b12_right.line.width = Pt(1)
    b12r_tf = b12_right.text_frame
    b12r_tf.word_wrap = False
    b12r_p = b12r_tf.paragraphs[0]
    b12r_p.alignment = PP_ALIGN.CENTER
    b12r_r = b12r_p.add_run()
    b12r_r.text = "💬 SIMULADOR DE CHAT"
    b12r_r.font.name = "Montserrat"
    b12r_r.font.size = Pt(11)
    b12r_r.font.bold = True
    b12r_r.font.color.rgb = RGBColor(186, 230, 253)

    tb12_r = slide3.shapes.add_textbox(Inches(7.153), Inches(2.46), Inches(5.23), Inches(3.60))
    tf12_r = tb12_r.text_frame
    tf12_r.word_wrap = True
    tf12_r.margin_left = tf12_r.margin_top = tf12_r.margin_right = tf12_r.margin_bottom = 0

    p12r_h = tf12_r.paragraphs[0]
    p12r_h.text = "Previsualización en Vivo"
    p12r_h.font.name = "Montserrat"
    p12r_h.font.size = Pt(17)
    p12r_h.font.bold = True
    p12r_h.font.color.rgb = RGBColor(56, 189, 248)
    p12r_h.space_after = Pt(3)

    p12r_sub = tf12_r.add_paragraph()
    p12r_sub.text = "Mira cómo cobra vida tu engaño en tiempo real antes de enviarlo al sistema."
    p12r_sub.font.name = "Inter"
    p12r_sub.font.size = Pt(11)
    p12r_sub.font.color.rgb = RGBColor(148, 163, 184)
    p12r_sub.space_after = Pt(12)

    items_12r = [
        ("👁️ Interfaz Real de WhatsApp: ", "Tu mensaje se construye en vivo con globos de chat y hora de envío."),
        ("🧠 Astucia vs. Detección: ", "¿Podrás crear un mensaje convincente sin dejar pistas demasiado obvias?"),
        ("💾 Guardado Automático: ", "Si se corta la conexión del aula, tu progreso se preserva intacto."),
        ("⏱️ Cronómetro en Vivo: ", "Sigue el tiempo del reloj sincronizado con la pantalla del aula.")
    ]

    for title_txt, body_txt in items_12r:
        pi = tf12_r.add_paragraph()
        pi.space_after = Pt(8)
        
        rt = pi.add_run()
        rt.text = title_txt
        rt.font.name = "Inter"
        rt.font.size = Pt(11.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(224, 242, 254)

        rb = pi.add_run()
        rb.text = body_txt
        rb.font.name = "Inter"
        rb.font.size = Pt(11)
        rb.font.color.rgb = RGBColor(148, 163, 184)

    callout12_r = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.153), Inches(6.20), Inches(5.23), Inches(0.56))
    callout12_r.fill.solid()
    callout12_r.fill.fore_color.rgb = RGBColor(12, 26, 46)
    callout12_r.line.color.rgb = RGBColor(3, 105, 161)
    callout12_r.line.width = Pt(1)
    c12r_tf = callout12_r.text_frame
    c12r_tf.word_wrap = True
    c12r_tf.margin_top = Inches(0.08)
    c12r_p = c12r_tf.paragraphs[0]
    c12r_p.alignment = PP_ALIGN.CENTER
    c12r_r1 = c12r_p.add_run()
    c12r_r1.text = "⚡ 100% Web: "
    c12r_r1.font.name = "Inter"
    c12r_r1.font.size = Pt(10.5)
    c12r_r1.font.bold = True
    c12r_r1.font.color.rgb = RGBColor(56, 189, 248)
    c12r_r2 = c12r_p.add_run()
    c12r_r2.text = "Sin registro ni instalación. Ingreso con QR gigante."
    c12r_r2.font.name = "Inter"
    c12r_r2.font.size = Pt(10)
    c12r_r2.font.color.rgb = RGBColor(226, 232, 240)


    # =========================================================================
    # SLIDE 13: Dinámica de Retroalimentación: Personal & Global
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    slide4.shapes.add_picture(io.BytesIO(bg_blob), 0, 0, Inches(13.333), Inches(7.50))

    tb_title4 = slide4.shapes.add_textbox(Inches(0.625), Inches(0.48), Inches(12.08), Inches(0.55))
    tf4 = tb_title4.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_top = tf4.margin_right = tf4.margin_bottom = 0
    p4 = tf4.paragraphs[0]
    p4.text = "Retroalimentación: Personal y Global"
    p4.font.name = "Montserrat"
    p4.font.size = Pt(27)
    p4.font.bold = True
    p4.font.color.rgb = RGBColor(248, 250, 252)

    tb_sub4 = slide4.shapes.add_textbox(Inches(0.625), Inches(1.12), Inches(12.08), Inches(0.45))
    tfs4 = tb_sub4.text_frame
    tfs4.word_wrap = True
    tfs4.margin_left = tfs4.margin_top = tfs4.margin_right = tfs4.margin_bottom = 0
    ps4 = tfs4.paragraphs[0]
    
    r4_1 = ps4.add_run()
    r4_1.text = "Evaluación instantánea en tu celular y "
    r4_1.font.name = "Inter"
    r4_1.font.size = Pt(13)
    r4_1.font.color.rgb = RGBColor(148, 163, 184)
    
    r4_2 = ps4.add_run()
    r4_2.text = "puesta en común grupal en pantalla gigante"
    r4_2.font.name = "Inter"
    r4_2.font.size = Pt(13)
    r4_2.font.bold = True
    r4_2.font.color.rgb = RGBColor(56, 189, 248)

    r4_3 = ps4.add_run()
    r4_3.text = "."
    r4_3.font.name = "Inter"
    r4_3.font.size = Pt(13)
    r4_3.font.color.rgb = RGBColor(148, 163, 184)

    # Left Card: Feedback Personal
    c13_left = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.625), card_y2, card_w2, card_h2)
    c13_left.fill.solid()
    c13_left.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c13_left.line.color.rgb = RGBColor(16, 185, 129)
    c13_left.line.width = Pt(1.5)

    b13_left = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.98), Inches(3.2), Inches(0.36))
    b13_left.fill.solid()
    b13_left.fill.fore_color.rgb = RGBColor(6, 78, 59)
    b13_left.line.color.rgb = RGBColor(16, 185, 129)
    b13_left.line.width = Pt(1)
    b13l_tf = b13_left.text_frame
    b13l_tf.word_wrap = False
    b13l_p = b13l_tf.paragraphs[0]
    b13l_p.alignment = PP_ALIGN.CENTER
    b13l_r = b13l_p.add_run()
    b13l_r.text = "📱 FEEDBACK PERSONAL"
    b13l_r.font.name = "Montserrat"
    b13l_r.font.size = Pt(11)
    b13l_r.font.bold = True
    b13l_r.font.color.rgb = RGBColor(167, 243, 208)

    tb13_l = slide4.shapes.add_textbox(Inches(0.95), Inches(2.46), Inches(5.23), Inches(3.60))
    tf13_l = tb13_l.text_frame
    tf13_l.word_wrap = True
    tf13_l.margin_left = tf13_l.margin_top = tf13_l.margin_right = tf13_l.margin_bottom = 0

    p13l_h = tf13_l.paragraphs[0]
    p13l_h.text = "Evaluación en tu Dispositivo"
    p13l_h.font.name = "Montserrat"
    p13l_h.font.size = Pt(17)
    p13l_h.font.bold = True
    p13l_h.font.color.rgb = RGBColor(52, 211, 153)
    p13l_h.space_after = Pt(3)

    p13l_sub = tf13_l.add_paragraph()
    p13l_sub.text = "Análisis técnico automatizado e instantáneo al momento de enviar."
    p13l_sub.font.name = "Inter"
    p13l_sub.font.size = Pt(11)
    p13l_sub.font.color.rgb = RGBColor(148, 163, 184)
    p13l_sub.space_after = Pt(12)

    items_13l = [
        ("📊 Puntaje de Efectividad: ", "Métricas de persuasión e impacto según las opciones elegidas."),
        ("🔍 Detección Forense: ", "Desglose punto por punto de qué pistas o errores delataron el engaño."),
        ("💡 Por Qué Funciona: ", "Explicación psicológica de por qué una víctima caería ante tu mensaje."),
        ("🛡️ Consejos de Blindaje: ", "Recomendaciones personalizadas para no ser engañado en la vida real.")
    ]

    for title_txt, body_txt in items_13l:
        pi = tf13_l.add_paragraph()
        pi.space_after = Pt(8)
        
        rt = pi.add_run()
        rt.text = title_txt
        rt.font.name = "Inter"
        rt.font.size = Pt(11.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(209, 250, 229)

        rb = pi.add_run()
        rb.text = body_txt
        rb.font.name = "Inter"
        rb.font.size = Pt(11)
        rb.font.color.rgb = RGBColor(148, 163, 184)

    callout13_l = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(6.20), Inches(5.23), Inches(0.56))
    callout13_l.fill.solid()
    callout13_l.fill.fore_color.rgb = RGBColor(10, 32, 24)
    callout13_l.line.color.rgb = RGBColor(4, 120, 87)
    callout13_l.line.width = Pt(1)
    c13l_tf = callout13_l.text_frame
    c13l_tf.word_wrap = True
    c13l_tf.margin_top = Inches(0.08)
    c13l_p = c13l_tf.paragraphs[0]
    c13l_p.alignment = PP_ALIGN.CENTER
    c13l_r1 = c13l_p.add_run()
    c13l_r1.text = "📈 Diagnóstico Instantáneo: "
    c13l_r1.font.name = "Inter"
    c13l_r1.font.size = Pt(10.5)
    c13l_r1.font.bold = True
    c13l_r1.font.color.rgb = RGBColor(52, 211, 153)
    c13l_r2 = c13l_p.add_run()
    c13l_r2.text = "Aprende de tus aciertos y descuidos."
    c13l_r2.font.name = "Inter"
    c13l_r2.font.size = Pt(10)
    c13l_r2.font.color.rgb = RGBColor(226, 232, 240)


    # Right Card: Feedback Global
    c13_right = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.828), card_y2, card_w2, card_h2)
    c13_right.fill.solid()
    c13_right.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c13_right.line.color.rgb = RGBColor(245, 158, 11)
    c13_right.line.width = Pt(1.5)

    b13_right = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.153), Inches(1.98), Inches(3.2), Inches(0.36))
    b13_right.fill.solid()
    b13_right.fill.fore_color.rgb = RGBColor(120, 53, 15)
    b13_right.line.color.rgb = RGBColor(245, 158, 11)
    b13_right.line.width = Pt(1)
    b13r_tf = b13_right.text_frame
    b13r_tf.word_wrap = False
    b13r_p = b13r_tf.paragraphs[0]
    b13r_p.alignment = PP_ALIGN.CENTER
    b13r_r = b13r_p.add_run()
    b13r_r.text = "📽️ PUESTA EN COMÚN"
    b13r_r.font.name = "Montserrat"
    b13r_r.font.size = Pt(11)
    b13r_r.font.bold = True
    b13r_r.font.color.rgb = RGBColor(253, 230, 138)

    tb13_r = slide4.shapes.add_textbox(Inches(7.153), Inches(2.46), Inches(5.23), Inches(3.60))
    tf13_r = tb13_r.text_frame
    tf13_r.word_wrap = True
    tf13_r.margin_left = tf13_r.margin_top = tf13_r.margin_right = tf13_r.margin_bottom = 0

    p13r_h = tf13_r.paragraphs[0]
    p13r_h.text = "Muro y Debate del Aula"
    p13r_h.font.name = "Montserrat"
    p13r_h.font.size = Pt(17)
    p13r_h.font.bold = True
    p13r_h.font.color.rgb = RGBColor(251, 191, 36)
    p13r_h.space_after = Pt(3)

    p13r_sub = tf13_r.add_paragraph()
    p13r_sub.text = "Proyección sincrónica en el proyector conducida por el facilitador."
    p13r_sub.font.name = "Inter"
    p13r_sub.font.size = Pt(11)
    p13r_sub.font.color.rgb = RGBColor(148, 163, 184)
    p13r_sub.space_after = Pt(12)

    items_13r = [
        ("📽️ Muro de Mensajes en Vivo: ", "Visualización interactiva de todos los mensajes armados por el curso."),
        ("📈 Estadísticas del Aula: ", "Descubrimos cuáles fueron los ganchos y anzuelos más elegidos."),
        ("🧠 Debate Colectivo: ", "Analizamos en grupo por qué caemos y cómo detectar las trampas a tiempo."),
        ("🏆 Podio & Cierre Lúdico: ", "Reconocimiento a los mensajes más convincentes y mejor blindados.")
    ]

    for title_txt, body_txt in items_13r:
        pi = tf13_r.add_paragraph()
        pi.space_after = Pt(8)
        
        rt = pi.add_run()
        rt.text = title_txt
        rt.font.name = "Inter"
        rt.font.size = Pt(11.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(254, 243, 199)

        rb = pi.add_run()
        rb.text = body_txt
        rb.font.name = "Inter"
        rb.font.size = Pt(11)
        rb.font.color.rgb = RGBColor(148, 163, 184)

    callout13_r = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.153), Inches(6.20), Inches(5.23), Inches(0.56))
    callout13_r.fill.solid()
    callout13_r.fill.fore_color.rgb = RGBColor(32, 24, 12)
    callout13_r.line.color.rgb = RGBColor(180, 83, 9)
    callout13_r.line.width = Pt(1)
    c13r_tf = callout13_r.text_frame
    c13r_tf.word_wrap = True
    c13r_tf.margin_top = Inches(0.08)
    c13r_p = c13r_tf.paragraphs[0]
    c13r_p.alignment = PP_ALIGN.CENTER
    c13r_r1 = c13r_p.add_run()
    c13r_r1.text = "🎉 Reflexión Grupal: "
    c13r_r1.font.name = "Inter"
    c13r_r1.font.size = Pt(10.5)
    c13r_r1.font.bold = True
    c13r_r1.font.color.rgb = RGBColor(251, 191, 36)
    c13r_r2 = c13r_p.add_run()
    c13r_r2.text = "Puntos en común y aprendizajes compartidos."
    c13r_r2.font.name = "Inter"
    c13r_r2.font.size = Pt(10)
    c13r_r2.font.color.rgb = RGBColor(226, 232, 240)


    # Reorder slides so the 4 new slides are at indices 9, 10, 11, 12 (slides 10, 11, 12, 13)
    # Original slides: 11
    # Added 4 slides -> total 15
    # The new slides are at indices 11, 12, 13, 14
    sldIdLst = prs.slides._sldIdLst
    s1_elem = sldIdLst[11]
    s2_elem = sldIdLst[12]
    s3_elem = sldIdLst[13]
    s4_elem = sldIdLst[14]
    
    sldIdLst.remove(s1_elem)
    sldIdLst.remove(s2_elem)
    sldIdLst.remove(s3_elem)
    sldIdLst.remove(s4_elem)
    
    # Insert at index 9 (after slide 9: Respuesta ante incidentes)
    sldIdLst.insert(9, s1_elem)
    sldIdLst.insert(10, s2_elem)
    sldIdLst.insert(11, s3_elem)
    sldIdLst.insert(12, s4_elem)

    out_path = "Ciudadanía Digital Segura - Taller Individual WhatsApp.pptx"
    prs.save(out_path)
    print(f"Presentation saved successfully with {len(prs.slides)} slides to '{out_path}'")

if __name__ == "__main__":
    build_presentation()
