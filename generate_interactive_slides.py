import os
import io
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_presentation():
    src_path = "Ciudadanía Digital Segura_backup.pptx"
    prs = pptx.Presentation(src_path)

    bg_blob = prs.slides[0].shapes[0].image.blob
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # --- SLIDE 10: Desafío Práctico: CSI Digital ---
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    slide1.shapes.add_picture(io.BytesIO(bg_blob), 0, 0, Inches(13.333), Inches(7.50))

    # Title
    tb_title = slide1.shapes.add_textbox(Inches(0.625), Inches(0.48), Inches(12.08), Inches(0.55))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = "Desafío Práctico: CSI Digital"
    p_title.font.name = "Montserrat"
    p_title.font.size = Pt(27)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(248, 250, 252)

    # Subtitle
    tb_sub = slide1.shapes.add_textbox(Inches(0.625), Inches(1.12), Inches(12.08), Inches(0.45))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    
    r1 = p_sub.add_run()
    r1.text = "Pondremos a prueba lo aprendido en una simulación interactiva en tiempo real: "
    r1.font.name = "Inter"
    r1.font.size = Pt(13)
    r1.font.color.rgb = RGBColor(148, 163, 184)
    
    r2 = p_sub.add_run()
    r2.text = "Ataque vs. Defensa Digital"
    r2.font.name = "Inter"
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(56, 189, 248)

    r3 = p_sub.add_run()
    r3.text = "."
    r3.font.name = "Inter"
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGBColor(148, 163, 184)

    # Cards Geometry
    card_w = Inches(5.88)
    card_h = Inches(5.30)
    card_y = Inches(1.70)
    
    # -------------------------------------------------------------
    # Left Card: RED TEAM
    # -------------------------------------------------------------
    c1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.625), card_y, card_w, card_h)
    c1.fill.solid()
    c1.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c1.line.color.rgb = RGBColor(239, 68, 68)
    c1.line.width = Pt(1.5)

    # Left Badge
    b1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.98), Inches(3.2), Inches(0.36))
    b1.fill.solid()
    b1.fill.fore_color.rgb = RGBColor(69, 10, 10)
    b1.line.color.rgb = RGBColor(239, 68, 68)
    b1.line.width = Pt(1)
    b1_tf = b1.text_frame
    b1_tf.word_wrap = False
    b1_p = b1_tf.paragraphs[0]
    b1_p.alignment = PP_ALIGN.CENTER
    b1_r = b1_p.add_run()
    b1_r.text = "🔴 RED TEAM  |  ATACANTES"
    b1_r.font.name = "Montserrat"
    b1_r.font.size = Pt(11)
    b1_r.font.bold = True
    b1_r.font.color.rgb = RGBColor(254, 202, 202)

    # Left Text Content
    tb_c1 = slide1.shapes.add_textbox(Inches(0.95), Inches(2.46), Inches(5.23), Inches(3.60))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_top = tf_c1.margin_right = tf_c1.margin_bottom = 0

    p_c1_h = tf_c1.paragraphs[0]
    p_c1_h.text = "Misión de Ingeniería Social"
    p_c1_h.font.name = "Montserrat"
    p_c1_h.font.size = Pt(17)
    p_c1_h.font.bold = True
    p_c1_h.font.color.rgb = RGBColor(248, 113, 113)
    p_c1_h.space_after = Pt(3)

    p_c1_sub = tf_c1.add_paragraph()
    p_c1_sub.text = "Objetivo: Ponerse en los zapatos del ciberdelincuente para comprender los mecanismos del engaño."
    p_c1_sub.font.name = "Inter"
    p_c1_sub.font.size = Pt(11)
    p_c1_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_c1_sub.space_after = Pt(12)

    items_red = [
        ("🎯 Analizar el Escenario: ", "Estudiar a la víctima potencial e identificar vulnerabilidades humanas y descuidos."),
        ("⚡ Diseñar el Pretexto: ", "Construir la narrativa de persuasión apelando a la urgencia (24h), el miedo o la curiosidad."),
        ("🪤 Trampa Técnica: ", "Definir el anzuelo exacto: remitente falseado, enlace fraudulento (.xyz) o adjunto malicioso.")
    ]

    for title_txt, body_txt in items_red:
        p_item = tf_c1.add_paragraph()
        p_item.space_after = Pt(8)
        
        rt = p_item.add_run()
        rt.text = title_txt
        rt.font.name = "Inter"
        rt.font.size = Pt(11.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(254, 226, 226)

        rb = p_item.add_run()
        rb.text = body_txt
        rb.font.name = "Inter"
        rb.font.size = Pt(11)
        rb.font.color.rgb = RGBColor(148, 163, 184)

    # Left Callout Box at bottom
    callout1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(6.20), Inches(5.23), Inches(0.56))
    callout1.fill.solid()
    callout1.fill.fore_color.rgb = RGBColor(24, 18, 30)
    callout1.line.color.rgb = RGBColor(153, 27, 27)
    callout1.line.width = Pt(1)
    c1_tf = callout1.text_frame
    c1_tf.word_wrap = True
    c1_tf.margin_top = Inches(0.08)
    c1_p = c1_tf.paragraphs[0]
    c1_p.alignment = PP_ALIGN.CENTER
    c1_r1 = c1_p.add_run()
    c1_r1.text = "📱 En tu Celular: "
    c1_r1.font.name = "Inter"
    c1_r1.font.size = Pt(10.5)
    c1_r1.font.bold = True
    c1_r1.font.color.rgb = RGBColor(248, 113, 113)
    c1_r2 = c1_p.add_run()
    c1_r2.text = "Redactar la estrategia del engaño y las trampas."
    c1_r2.font.name = "Inter"
    c1_r2.font.size = Pt(10)
    c1_r2.font.color.rgb = RGBColor(226, 232, 240)


    # -------------------------------------------------------------
    # Right Card: BLUE TEAM
    # -------------------------------------------------------------
    c2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.828), card_y, card_w, card_h)
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c2.line.color.rgb = RGBColor(14, 165, 233)
    c2.line.width = Pt(1.5)

    # Right Badge
    b2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.153), Inches(1.98), Inches(3.2), Inches(0.36))
    b2.fill.solid()
    b2.fill.fore_color.rgb = RGBColor(8, 47, 73)
    b2.line.color.rgb = RGBColor(14, 165, 233)
    b2.line.width = Pt(1)
    b2_tf = b2.text_frame
    b2_tf.word_wrap = False
    b2_p = b2_tf.paragraphs[0]
    b2_p.alignment = PP_ALIGN.CENTER
    b2_r = b2_p.add_run()
    b2_r.text = "🔵 BLUE TEAM  |  DEFENSORES"
    b2_r.font.name = "Montserrat"
    b2_r.font.size = Pt(11)
    b2_r.font.bold = True
    b2_r.font.color.rgb = RGBColor(186, 230, 253)

    # Right Text Content
    tb_c2 = slide1.shapes.add_textbox(Inches(7.153), Inches(2.46), Inches(5.23), Inches(3.60))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = tf_c2.margin_top = tf_c2.margin_right = tf_c2.margin_bottom = 0

    p_c2_h = tf_c2.paragraphs[0]
    p_c2_h.text = "Misión de Análisis Forense"
    p_c2_h.font.name = "Montserrat"
    p_c2_h.font.size = Pt(17)
    p_c2_h.font.bold = True
    p_c2_h.font.color.rgb = RGBColor(56, 189, 248)
    p_c2_h.space_after = Pt(3)

    p_c2_sub = tf_c2.add_paragraph()
    p_c2_sub.text = "Objetivo: Inspeccionar como analistas de ciberseguridad (SOC) para neutralizar la amenaza."
    p_c2_sub.font.name = "Inter"
    p_c2_sub.font.size = Pt(11)
    p_c2_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_c2_sub.space_after = Pt(12)

    items_blue = [
        ("🔍 Inspección Técnica: ", "Auditar remitentes reales, dominios sospechosos, enlaces engañosos y archivos."),
        ("🛡️ Detectar 5 Anomalías: ", "Marcar en el checklist interactivo las 5 señales forenses ocultas en el caso."),
        ("📋 Protocolo de Defensa: ", "Redactar 3 recomendaciones prácticas y prioritarias para alertar a la comunidad.")
    ]

    for title_txt, body_txt in items_blue:
        p_item = tf_c2.add_paragraph()
        p_item.space_after = Pt(8)
        
        rt = p_item.add_run()
        rt.text = title_txt
        rt.font.name = "Inter"
        rt.font.size = Pt(11.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(224, 242, 254)

        rb = p_item.add_run()
        rb.text = body_txt
        rb.font.name = "Inter"
        rb.font.size = Pt(11)
        rb.font.color.rgb = RGBColor(148, 163, 184)

    # Right Callout Box at bottom
    callout2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.153), Inches(6.20), Inches(5.23), Inches(0.56))
    callout2.fill.solid()
    callout2.fill.fore_color.rgb = RGBColor(12, 26, 46)
    callout2.line.color.rgb = RGBColor(3, 105, 161)
    callout2.line.width = Pt(1)
    c2_tf = callout2.text_frame
    c2_tf.word_wrap = True
    c2_tf.margin_top = Inches(0.08)
    c2_p = c2_tf.paragraphs[0]
    c2_p.alignment = PP_ALIGN.CENTER
    c2_r1 = c2_p.add_run()
    c2_r1.text = "📱 En tu Celular: "
    c2_r1.font.name = "Inter"
    c2_r1.font.size = Pt(10.5)
    c2_r1.font.bold = True
    c2_r1.font.color.rgb = RGBColor(56, 189, 248)
    c2_r2 = c2_p.add_run()
    c2_r2.text = "Marcar 5 anomalías en el checklist + 3 defensas."
    c2_r2.font.name = "Inter"
    c2_r2.font.size = Pt(10)
    c2_r2.font.color.rgb = RGBColor(226, 232, 240)


    # =========================================================================
    # --- SLIDE 11: ¿Cómo Participamos? — Dinámica en Vivo ---
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.shapes.add_picture(io.BytesIO(bg_blob), 0, 0, Inches(13.333), Inches(7.50))

    # Title
    tb_title2 = slide2.shapes.add_textbox(Inches(0.625), Inches(0.48), Inches(12.08), Inches(0.55))
    tf_title2 = tb_title2.text_frame
    tf_title2.word_wrap = True
    tf_title2.margin_left = tf_title2.margin_top = tf_title2.margin_right = tf_title2.margin_bottom = 0
    p_title2 = tf_title2.paragraphs[0]
    p_title2.text = "¿Cómo Participamos? — Dinámica en Vivo"
    p_title2.font.name = "Montserrat"
    p_title2.font.size = Pt(27)
    p_title2.font.bold = True
    p_title2.font.color.rgb = RGBColor(248, 250, 252)

    # Subtitle
    tb_sub2 = slide2.shapes.add_textbox(Inches(0.625), Inches(1.12), Inches(12.08), Inches(0.45))
    tf_sub2 = tb_sub2.text_frame
    tf_sub2.word_wrap = True
    tf_sub2.margin_left = tf_sub2.margin_top = tf_sub2.margin_right = tf_sub2.margin_bottom = 0
    p_sub2 = tf_sub2.paragraphs[0]
    
    r2_1 = p_sub2.add_run()
    r2_1.text = "Plataforma 100% web desde tu celular, sin descargas ni registros. "
    r2_1.font.name = "Inter"
    r2_1.font.size = Pt(13)
    r2_1.font.color.rgb = RGBColor(148, 163, 184)
    
    r2_2 = p_sub2.add_run()
    r2_2.text = "¡Misión contrarreloj por equipos!"
    r2_2.font.name = "Inter"
    r2_2.font.size = Pt(13)
    r2_2.font.bold = True
    r2_2.font.color.rgb = RGBColor(56, 189, 248)

    # 3 Cards Layout
    c3_w = Inches(3.82)
    c3_h = Inches(5.30)
    c3_y = Inches(1.70)
    
    cards_data = [
        {
            "left": Inches(0.625),
            "border_color": RGBColor(14, 165, 233),
            "badge_bg": RGBColor(8, 47, 73),
            "badge_txt": "PASO 1",
            "badge_color": RGBColor(186, 230, 253),
            "title": "1. Ingreso al Juego",
            "title_color": RGBColor(56, 189, 248),
            "subtitle": "Conexión instantánea en segundos desde cualquier celular.",
            "items": [
                ("📷 Escaneo de QR: ", "Apunta la cámara de tu celular al código QR gigante en la proyección."),
                ("🔑 Código PIN: ", "O ingresa a la web del juego e introduce el código de sesión de 4 dígitos."),
                ("👥 Elige tu Equipo: ", "Selecciona tu grupo y descubre tu rol asignado (Red o Blue Team).")
            ],
            "callout_bg": RGBColor(12, 26, 46),
            "callout_border": RGBColor(3, 105, 161),
            "callout_text": "⚡ 100% web, sin instalar nada"
        },
        {
            "left": Inches(4.756),
            "border_color": RGBColor(245, 158, 11),
            "badge_bg": RGBColor(120, 53, 15),
            "badge_txt": "PASO 2",
            "badge_color": RGBColor(253, 230, 138),
            "title": "2. Misión en Vivo",
            "title_color": RGBColor(251, 191, 36),
            "subtitle": "Análisis en equipo sincronizado con el cronómetro del aula.",
            "items": [
                ("⏳ Tiempo Compartido: ", "El reloj corre sincronizado en la pantalla gigante y en los teléfonos."),
                ("💬 Debate en Grupo: ", "Analicen el caso asignado (correo bancario o chat de WhatsApp sospechoso)."),
                ("🚀 Envío de Misión: ", "Completen el formulario y envíen antes del aviso final de 2 minutos.")
            ],
            "callout_bg": RGBColor(32, 24, 12),
            "callout_border": RGBColor(180, 83, 9),
            "callout_text": "💾 Guardado automático de respuestas"
        },
        {
            "left": Inches(8.888),
            "border_color": RGBColor(16, 185, 129),
            "badge_bg": RGBColor(6, 78, 59),
            "badge_txt": "PASO 3",
            "badge_color": RGBColor(167, 243, 208),
            "title": "3. Revelación y Podio",
            "title_color": RGBColor(52, 211, 153),
            "subtitle": "Puesta en común, análisis forense y ranking lúdico final.",
            "items": [
                ("📽️ Puesta en Común: ", "Proyección en grande de las respuestas enviadas por cada equipo."),
                ("🎯 Revelación Forense: ", "Descubrimos juntos las 5 anomalías reales con alertas animadas."),
                ("🏆 Podio y Puntos: ", "Tabla de posiciones con puntos por aciertos y bonus de velocidad.")
            ],
            "callout_bg": RGBColor(10, 32, 24),
            "callout_border": RGBColor(4, 120, 87),
            "callout_text": "🎉 ¡Cálculo de puntos y confeti digital!"
        }
    ]

    for cdata in cards_data:
        # Container
        card_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cdata["left"], c3_y, c3_w, c3_h)
        card_box.fill.solid()
        card_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
        card_box.line.color.rgb = cdata["border_color"]
        card_box.line.width = Pt(1.5)

        # Step Badge
        badge_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cdata["left"] + Inches(0.25), Inches(1.98), Inches(1.2), Inches(0.35))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = cdata["badge_bg"]
        badge_box.line.color.rgb = cdata["border_color"]
        badge_box.line.width = Pt(1)
        b_tf = badge_box.text_frame
        b_tf.word_wrap = False
        b_p = b_tf.paragraphs[0]
        b_p.alignment = PP_ALIGN.CENTER
        b_r = b_p.add_run()
        b_r.text = cdata["badge_txt"]
        b_r.font.name = "Montserrat"
        b_r.font.size = Pt(10.5)
        b_r.font.bold = True
        b_r.font.color.rgb = cdata["badge_color"]

        # Content
        tb = slide2.shapes.add_textbox(cdata["left"] + Inches(0.25), Inches(2.46), c3_w - Inches(0.5), Inches(3.60))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_h = tf.paragraphs[0]
        p_h.text = cdata["title"]
        p_h.font.name = "Montserrat"
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = cdata["title_color"]
        p_h.space_after = Pt(3)

        p_sub = tf.add_paragraph()
        p_sub.text = cdata["subtitle"]
        p_sub.font.name = "Inter"
        p_sub.font.size = Pt(10.5)
        p_sub.font.color.rgb = RGBColor(148, 163, 184)
        p_sub.space_after = Pt(12)

        for ititle, ibody in cdata["items"]:
            p_item = tf.add_paragraph()
            p_item.space_after = Pt(8)
            
            rt = p_item.add_run()
            rt.text = ititle
            rt.font.name = "Inter"
            rt.font.size = Pt(11)
            rt.font.bold = True
            rt.font.color.rgb = RGBColor(241, 245, 249)

            rb = p_item.add_run()
            rb.text = ibody
            rb.font.name = "Inter"
            rb.font.size = Pt(10.5)
            rb.font.color.rgb = RGBColor(148, 163, 184)

        # Callout Pill at bottom
        callout_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cdata["left"] + Inches(0.25), Inches(6.20), c3_w - Inches(0.5), Inches(0.56))
        callout_box.fill.solid()
        callout_box.fill.fore_color.rgb = cdata["callout_bg"]
        callout_box.line.color.rgb = cdata["callout_border"]
        callout_box.line.width = Pt(1)
        co_tf = callout_box.text_frame
        co_tf.word_wrap = True
        co_tf.margin_top = Inches(0.08)
        co_p = co_tf.paragraphs[0]
        co_p.alignment = PP_ALIGN.CENTER
        co_r = co_p.add_run()
        co_r.text = cdata["callout_text"]
        co_r.font.name = "Inter"
        co_r.font.size = Pt(10)
        co_r.font.bold = True
        co_r.font.color.rgb = cdata["badge_color"]

    # Reorder slides so the 2 new slides are at positions 10 and 11 (indices 9 and 10)
    sldIdLst = prs.slides._sldIdLst
    slide1_elem = sldIdLst[11]
    slide2_elem = sldIdLst[12]
    
    sldIdLst.remove(slide1_elem)
    sldIdLst.remove(slide2_elem)
    
    sldIdLst.insert(9, slide1_elem)
    sldIdLst.insert(10, slide2_elem)

    out_path = "Ciudadanía Digital Segura.pptx"
    prs.save(out_path)
    print(f"Presentation saved successfully with {len(prs.slides)} slides to {out_path}")

if __name__ == "__main__":
    build_presentation()
