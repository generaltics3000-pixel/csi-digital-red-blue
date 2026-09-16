import pptx

with open("/home/dario/.gemini/antigravity-ide/brain/b9af9946-92a1-4c0c-9a97-0b6a6e1b42f5/scratch/border_orig.png", "rb") as f:
    border_blob = f.read()

with open("/home/dario/.gemini/antigravity-ide/brain/b9af9946-92a1-4c0c-9a97-0b6a6e1b42f5/scratch/slide3_rounded.png", "rb") as f:
    s3_blob = f.read()

with open("/home/dario/.gemini/antigravity-ide/brain/b9af9946-92a1-4c0c-9a97-0b6a6e1b42f5/scratch/slide7_rounded.png", "rb") as f:
    s7_blob = f.read()

files = [
    "Ciudadanía Digital Segura.pptx",
    "Ciudadanía Digital Segura - Taller Individual WhatsApp.pptx"
]

for filename in files:
    prs = pptx.Presentation(filename)
    
    # Slide 3 (index 2)
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.name == "Google Shape;117;p15":
            blip = shape._element.xpath('.//a:blip')
            rId = blip[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            shape.part.related_part(rId)._blob = border_blob
        elif shape.name == "Google Shape;120;p15":
            blip = shape._element.xpath('.//a:blip')
            rId = blip[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            shape.part.related_part(rId)._blob = s3_blob

    # Slide 7 (index 6)
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.name == "Google Shape;196;p19":
            blip = shape._element.xpath('.//a:blip')
            rId = blip[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            shape.part.related_part(rId)._blob = border_blob
        elif shape.name == "Google Shape;199;p19":
            blip = shape._element.xpath('.//a:blip')
            rId = blip[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            shape.part.related_part(rId)._blob = s7_blob

    prs.save(filename)
    print(f"Fixed {filename} properly")

